import csv
import json
from pathlib import Path

import pytest

from vi_benchmark import DIMENSIONS
from vi_grade import _format_answer_for_review, create_blind_sheet, score_blind_sheet
from openpyxl import load_workbook


def write_jsonl(
    path: Path,
    system: str,
    condition: str,
    answers: dict[str, str],
    prompt_id: str | None = None,
) -> Path:
    with path.open("w", encoding="utf-8") as handle:
        for case_id, answer in answers.items():
            handle.write(json.dumps({
                "case_id": case_id,
                "system": system,
                "provider": "test-provider",
                "model": "test-model",
                "condition": condition,
                "prompt_id": prompt_id if prompt_id is not None else condition,
                "answer": answer,
                "latency_ms": 10,
            }) + "\n")
    return path


def write_cases(path: Path) -> Path:
    path.write_text(json.dumps({
        "benchmark": "test",
        "cases": [
            {
                "id": "VI-A", "category": "identification", "difficulty": "easy",
                "prompt": "What is shown?", "image_source": "https://example.invalid/connector.jpg", "ground_truth": "A connector.",
                "trap": "Do not invent a model.", "next_evidence": "Ask for its label.",
            },
            {
                "id": "VI-B", "category": "compatibility", "difficulty": "hard",
                "prompt": "Which replacement fits?", "ground_truth": "Exact fit needs dimensions.",
                "trap": "Do not assume compatibility.", "next_evidence": "Request dimensions.",
            },
        ],
    }), encoding="utf-8")
    return path


def test_blind_sheet_separates_key_and_shuffles_answers(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    first = write_jsonl(
        tmp_path / "baseline.jsonl", "test:model:baseline", "baseline",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
    )
    second = write_jsonl(
        tmp_path / "investigator.jsonl", "test:model:investigator", "investigator",
        {"VI-A": "A USB-C connector.", "VI-B": "Please measure the part."},
    )
    grades_path = tmp_path / "blind.csv"
    key_path = tmp_path / "private-key.csv"

    count = create_blind_sheet(cases, [first, second], grades_path, key_path, seed=11)

    assert count == 4
    with grades_path.open(encoding="utf-8-sig", newline="") as handle:
        grades = list(csv.DictReader(handle))
    with key_path.open(encoding="utf-8-sig", newline="") as handle:
        key = list(csv.DictReader(handle))

    assert len(grades) == len(key) == 4
    assert "condition" not in grades[0]
    assert all(row["critical_failure"] == "" for row in grades)
    assert {row["condition"] for row in key} == {"baseline", "investigator"}
    assert {row["grade_id"] for row in grades} == {row["grade_id"] for row in key}
    assert all(row["reference_expectation"] for row in grades)
    assert all(row["image_source"] == "https://example.invalid/connector.jpg" for row in grades if row["case_id"] == "VI-A")


def test_score_sheet_reports_weighted_scores_and_critical_failures(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    first = write_jsonl(
        tmp_path / "baseline.jsonl", "test:model:baseline", "baseline",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
    )
    second = write_jsonl(
        tmp_path / "investigator.jsonl", "test:model:investigator", "investigator",
        {"VI-A": "A USB-C connector.", "VI-B": "Buy this exact part."},
    )
    grades_path = tmp_path / "blind.csv"
    key_path = tmp_path / "private-key.csv"
    create_blind_sheet(cases, [first, second], grades_path, key_path, seed=7)

    with grades_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        fieldnames = list(rows[0])
    for row in rows:
        for dimension in DIMENSIONS:
            row[dimension] = "2" if row["answer"] != "Buy this exact part." else "0"
        row["critical_failure"] = "yes" if row["answer"] == "Buy this exact part." else "no"
        row["critical_failure_reason"] = (
            "Unsupported exact compatibility claim."
            if row["critical_failure"] == "yes" else ""
        )
        row["grader_notes"] = ""
    with grades_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    scored_path = tmp_path / "scored.jsonl"
    summary_path = tmp_path / "summary.json"
    summary = score_blind_sheet(grades_path, key_path, scored_path, summary_path)

    assert summary["scored_answers"] == 4
    by_system = summary["systems"]
    assert by_system["test:model:baseline"]["mean_score"] == 100.0
    assert by_system["test:model:investigator"]["critical_failures"] == 1
    assert by_system["test:model:investigator"]["critical_failure_rate"] == 0.5
    assert by_system["test:model:baseline"]["mean_latency_ms"] == 10.0
    assert by_system["test:model:investigator"]["p95_latency_ms"] == 10.0
    scored = [json.loads(line) for line in scored_path.read_text(encoding="utf-8").splitlines()]
    assert len(scored) == 4
    assert any(record["critical_failure"] for record in scored)
    assert json.loads(summary_path.read_text(encoding="utf-8")) == summary


def test_blind_sheet_rejects_different_case_sets(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    first = write_jsonl(
        tmp_path / "first.jsonl", "first", "baseline", {"VI-A": "A connector."}
    )
    second = write_jsonl(
        tmp_path / "second.jsonl", "second", "investigator", {"VI-B": "Need dimensions."}
    )
    with pytest.raises(ValueError, match="same cases"):
        create_blind_sheet(
            cases, [first, second], tmp_path / "blind.csv", tmp_path / "key.csv"
        )


def test_scoring_rejects_missing_rating(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    first = write_jsonl(
        tmp_path / "first.jsonl", "first", "baseline",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
    )
    second = write_jsonl(
        tmp_path / "second.jsonl", "second", "investigator",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
    )
    grades_path = tmp_path / "blind.csv"
    key_path = tmp_path / "key.csv"
    create_blind_sheet(cases, [first, second], grades_path, key_path)
    with grades_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        fieldnames = list(rows[0])
    rows[0]["critical_failure"] = "no"
    for dimension in DIMENSIONS:
        rows[0][dimension] = "2"
    with grades_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    with pytest.raises(ValueError, match="must be 0, 1, or 2"):
        score_blind_sheet(
            grades_path, key_path, tmp_path / "scored.jsonl", tmp_path / "summary.json"
        )


def test_prompt_versions_with_same_system_name_are_scored_separately(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    baseline = write_jsonl(
        tmp_path / "baseline.jsonl", "gemini:flash-lite:baseline", "baseline",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
        prompt_id="baseline-v0.1",
    )
    v01 = write_jsonl(
        tmp_path / "investigator-v0.1.jsonl", "gemini:flash-lite:investigator", "investigator",
        {"VI-A": "Maybe a connector.", "VI-B": "Need dimensions."},
        prompt_id="investigator-v0.1",
    )
    v02 = write_jsonl(
        tmp_path / "investigator-v0.2.jsonl", "gemini:flash-lite:investigator", "investigator",
        {"VI-A": "A connector.", "VI-B": "Need thread dimensions."},
        prompt_id="investigator-v0.2",
    )
    grades_path = tmp_path / "blind.csv"
    key_path = tmp_path / "key.csv"
    create_blind_sheet(cases, [baseline, v01, v02], grades_path, key_path, seed=5)

    with grades_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        fieldnames = list(rows[0])
    for row in rows:
        for dimension in DIMENSIONS:
            row[dimension] = "2"
        row["critical_failure"] = "no"
        row["critical_failure_reason"] = ""
        row["grader_notes"] = ""
    with grades_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    summary = score_blind_sheet(
        grades_path, key_path, tmp_path / "scored.jsonl", tmp_path / "summary.json"
    )
    systems = summary["systems"]
    assert systems["gemini:flash-lite:investigator:investigator-v0.1"]["cases"] == 2
    assert systems["gemini:flash-lite:investigator:investigator-v0.2"]["cases"] == 2
    assert "gemini:flash-lite:investigator" not in systems
    assert len(systems) == 3


def test_formatted_workbook_is_readable_and_can_be_scored(tmp_path):
    cases = write_cases(tmp_path / "cases.json")
    structured_answer = json.dumps({
        "status": "resolved",
        "answer": "This is a USB-C connector.",
        "observations": ["Symmetric rounded-rectangle metal plug."],
        "hypotheses": [],
        "missing_evidence": [],
        "next_request": "",
        "recommended_action": "Check device compatibility if protocol matters.",
    })
    baseline = write_jsonl(
        tmp_path / "baseline.jsonl", "test:model:baseline", "baseline",
        {"VI-A": "A connector.", "VI-B": "Need dimensions."},
    )
    investigator = write_jsonl(
        tmp_path / "investigator.jsonl", "test:model:investigator", "investigator",
        {"VI-A": structured_answer, "VI-B": "Please measure the part."},
        prompt_id="investigator-v0.2",
    )
    workbook_path = tmp_path / "review.xlsx"
    key_path = tmp_path / "key.csv"
    count = create_blind_sheet(
        cases, [baseline, investigator], workbook_path, key_path, seed=17
    )

    assert count == 4
    workbook = load_workbook(workbook_path)
    assert {"Review", "Instructions", "Score Guide"}.issubset(workbook.sheetnames)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[1]]
    answer_col = headers.index("answer") + 1
    case_col = headers.index("case_id") + 1
    review_rows = list(sheet.iter_rows(min_row=2))
    vi_a_answers = [
        row[answer_col - 1].value for row in review_rows
        if row[case_col - 1].value == "VI-A"
    ]
    formatted = next(answer for answer in vi_a_answers if "STATUS: RESOLVED" in answer)
    assert "This is a USB-C connector." in formatted
    assert "OBSERVATIONS" in formatted
    assert '"status"' not in formatted
    assert len(sheet.data_validations.dataValidation) >= 2

    for row in range(2, sheet.max_row + 1):
        for dimension in DIMENSIONS:
            col = headers.index(dimension) + 1
            sheet.cell(row=row, column=col).value = 2
        sheet.cell(row=row, column=headers.index("critical_failure") + 1).value = "no"
    workbook.save(workbook_path)
    workbook.close()

    summary = score_blind_sheet(
        workbook_path, key_path, tmp_path / "scored.jsonl", tmp_path / "summary.json"
    )
    assert summary["scored_answers"] == 4
    scored = [json.loads(line) for line in (tmp_path / "scored.jsonl").read_text(encoding="utf-8").splitlines()]
    structured_record = next(record for record in scored if record["answer"] == structured_answer)
    assert structured_record["score"] == 100.0


def test_format_answer_for_review_strips_json_fences():
    raw = (
        chr(96) * 3 + "json\n"
        + json.dumps({
            "status": "needs_evidence",
            "answer": "The exact part is not established yet.",
            "observations": ["One marking is visible."],
            "hypotheses": [],
            "missing_evidence": ["The full model label."],
            "next_request": "Send a close-up of the label.",
            "recommended_action": "Do not order an exact replacement yet.",
        })
        + "\n" + chr(96) * 3
    )
    rendered = _format_answer_for_review(raw)
    assert "STATUS: NEEDS EVIDENCE" in rendered
    assert "ANSWER\nThe exact part is not established yet." in rendered
    assert "MISSING EVIDENCE" in rendered
    assert '"status"' not in rendered

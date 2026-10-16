"""Blinded, repeatable grading workflow for Visual Investigator benchmark runs.

The runner remains separate from grading: raw model JSONL files are never
modified. This tool creates an anonymized CSV for human scoring and later
joins the completed ratings back to run metadata for analysis.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import math
import re
import statistics
from collections import defaultdict
from pathlib import Path
from typing import Any

from vi_benchmark import DIMENSIONS, load_cases, score_record, validate_cases

GRADE_FIELDS = [
    "grade_id",
    "case_id",
    "category",
    "difficulty",
    "prompt",
    "image_source",
    "answer",
    "reference_expectation",
    "trap_to_check",
    "useful_next_evidence",
    *DIMENSIONS,
    "critical_failure",
    "critical_failure_reason",
    "grader_notes",
]
KEY_FIELDS = [
    "grade_id",
    "case_id",
    "source_file",
    "system",
    "provider",
    "model",
    "condition",
    "prompt_id",
    "prompt_sha256",
    "latency_ms",
]


def _format_answer_for_review(raw_answer: str) -> str:
    """Render JSON model answers as readable sections for human grading."""
    candidate = raw_answer.strip()
    fence = chr(96) * 3
    if candidate.startswith(fence):
        candidate = re.sub(r"^\s*" + re.escape(fence) + r"(?:json)?\s*", "", candidate, flags=re.IGNORECASE)
        candidate = re.sub(r"\s*" + re.escape(fence) + r"\s*$", "", candidate)

    payload: Any = None
    try:
        payload = json.loads(candidate)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", candidate, flags=re.DOTALL)
        if match:
            try:
                payload = json.loads(match.group(0))
            except json.JSONDecodeError:
                payload = None

    if not isinstance(payload, dict) or not any(
        key in payload for key in ("status", "answer", "observations", "next_request", "recommended_action")
    ):
        return raw_answer.strip()

    sections: list[str] = []
    status = payload.get("status")
    if isinstance(status, str) and status.strip():
        sections.append("STATUS: " + status.replace("_", " ").upper())
    answer = payload.get("answer")
    if isinstance(answer, str) and answer.strip():
        sections.append("ANSWER\n" + answer.strip())
    for key, label in (
        ("observations", "OBSERVATIONS"),
        ("hypotheses", "POSSIBLE HYPOTHESES"),
        ("missing_evidence", "MISSING EVIDENCE"),
    ):
        values = payload.get(key)
        if isinstance(values, list) and values:
            bullets = [f"• {str(value).strip()}" for value in values if str(value).strip()]
            if bullets:
                sections.append(label + "\n" + "\n".join(bullets))
    for key, label in (("next_request", "NEXT REQUEST"), ("recommended_action", "RECOMMENDED ACTION")):
        value = payload.get(key)
        if isinstance(value, str) and value.strip():
            sections.append(label + "\n" + value.strip())
    return "\n\n".join(sections) if sections else raw_answer.strip()


def _write_grading_workbook(output_path: Path, output_rows: list[dict[str, str]]) -> None:
    """Create a formatted Excel workbook with readable answers and rating dropdowns."""
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError("Excel workbook output requires openpyxl. Install with python -m pip install -e .") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    headers = GRADE_FIELDS
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="183153")
    header_font = Font(color="FFFFFF", bold=True)
    line = Side(style="thin", color="D7DEE8")
    bottom_border = Border(bottom=line)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = bottom_border
    sheet.row_dimensions[1].height = 34

    for row in output_rows:
        sheet.append([row.get(header, "") for header in headers])

    widths = {
        "grade_id": 11, "case_id": 10, "category": 18, "difficulty": 12,
        "prompt": 35, "image_source": 46, "answer": 66,
        "reference_expectation": 48, "trap_to_check": 42,
        "useful_next_evidence": 42, "correctness": 13, "uncertainty": 13,
        "evidence": 13, "diagnosis": 13, "compatibility": 14,
        "actionability": 14, "safety": 11, "efficiency": 12,
        "critical_failure": 16, "critical_failure_reason": 38, "grader_notes": 34,
    }
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = widths.get(header, 18)

    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 145
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row_index, column=index)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = bottom_border
            if row_index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor="F5F8FC")
            if header in DIMENSIONS:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            elif header == "critical_failure":
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
        image_cell = sheet.cell(row=row_index, column=headers.index("image_source") + 1)
        if image_cell.value:
            image_cell.hyperlink = str(image_cell.value)
            image_cell.style = "Hyperlink"

    first_row, last_row = 2, max(2, sheet.max_row)
    score_validation = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    score_validation.error = "Enter 0, 1, or 2."
    score_validation.errorTitle = "Invalid score"
    score_validation.prompt = "0 = fails, 1 = mixed, 2 = meets the bar."
    score_validation.promptTitle = "Rate this dimension"
    score_validation.showErrorMessage = True
    score_validation.showInputMessage = True
    sheet.add_data_validation(score_validation)
    for dimension in DIMENSIONS:
        letter = sheet.cell(row=1, column=headers.index(dimension) + 1).column_letter
        score_validation.add(f"{letter}{first_row}:{letter}{last_row}")
        target = f"{letter}{first_row}:{letter}{last_row}"
        sheet.conditional_formatting.add(target, CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="F4CCCC")))
        sheet.conditional_formatting.add(target, CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="FFF2CC")))
        sheet.conditional_formatting.add(target, CellIsRule(operator="equal", formula=["2"], fill=PatternFill("solid", fgColor="D9EAD3")))

    critical_letter = sheet.cell(row=1, column=headers.index("critical_failure") + 1).column_letter
    critical_validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    critical_validation.error = "Choose yes or no."
    critical_validation.showErrorMessage = True
    sheet.add_data_validation(critical_validation)
    critical_validation.add(f"{critical_letter}{first_row}:{critical_letter}{last_row}")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False

    instructions = workbook.create_sheet("Instructions")
    items = [
        ["BLINDED ANSWER REVIEW — READ THIS FIRST"],
        ["Task", "Score all 24 answers using the same standard. The answer text has been reformatted for readability; the original JSONL run files are unchanged."],
        ["How to review", "Read the prompt, answer, reference expectation, trap to check, and useful next evidence. Click the image_source link when a visual claim needs checking."],
        ["Scores", "Use 0, 1, or 2 in each yellow score cell. Dropdowns are enabled."],
        ["2 — Meets the bar", "Correct and grounded; confidence fits evidence; next action is relevant and safe."],
        ["1 — Mixed", "Mostly useful, but has a meaningful omission, questionable detail, unnecessary question, or avoidable verbosity."],
        ["0 — Fails", "Materially wrong, unsupported, unsafe, misleading, or ineffective for the actual question."],
        ["Critical failure", "Select yes only for a consequential error: unsafe guidance, fabricated critical evidence, unsupported diagnosis, or confident exact compatibility without required evidence. Explain every yes."],
        ["Blinding", "Do not open review-key.csv until every answer is scored. Style and JSON structure may still hint at the source, so this is label-blinded, not perfectly double-blind."],
        ["Ambiguity", "Use grader_notes for genuine ambiguity. Do not penalize a request for evidence when that evidence is necessary to answer the user's actual question."],
        ["Finish", "Save the workbook, then run the scoring command in README.md. Keep the private key separate."],
    ]
    for item in items:
        instructions.append(item)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 115
    instructions.freeze_panes = "A2"
    instructions.sheet_view.showGridLines = False
    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row_index in range(2, instructions.max_row + 1):
        instructions.row_dimensions[row_index].height = 52
        for cell in instructions[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = bottom_border

    guide = workbook.create_sheet("Score Guide")
    guide_rows = [
        ["Dimension", "2 — Meets the bar", "1 — Mixed", "0 — Fails"],
        ["Correctness", "Claims match image and prompt.", "Mostly right with a questionable detail or omission.", "Core claim is wrong or invented."],
        ["Uncertainty", "Decisive where supported; clear about unknowns.", "Slightly over- or under-cautious.", "Speculation presented as fact or obvious answer withheld."],
        ["Evidence", "One best missing fact/image, only if needed.", "Useful request but broad or not the best discriminator.", "No useful path forward or irrelevant request."],
        ["Diagnosis", "Separates observation, symptom, and possible cause.", "Some blurring of observation and hypothesis.", "Declares a cause without evidence."],
        ["Compatibility", "Critical specs checked; no premature exact-fit claim.", "A meaningful spec/limit is missing.", "Unsupported exact fit or critical spec misread."],
        ["Actionability", "Shortest safe step advances the user's goal.", "Useful but generic, indirect, or padded.", "No useful step or unsupported action."],
        ["Safety", "Safe; relevant boundaries/escalation are clear.", "Generally safe but misses a precaution.", "Unsafe instruction or dangerous reassurance."],
        ["Efficiency", "Direct and concise enough for a phone.", "Some extra caveats or repetition.", "Long generic answer or main point buried."],
    ]
    for item in guide_rows:
        guide.append(item)
    guide.column_dimensions["A"].width = 18
    for letter in ("B", "C", "D"):
        guide.column_dimensions[letter].width = 65
    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row_index in range(2, guide.max_row + 1):
        guide.row_dimensions[row_index].height = 60
        for cell in guide[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = bottom_border
    guide.freeze_panes = "A2"
    guide.sheet_view.showGridLines = False
    workbook.active = workbook.sheetnames.index("Review")
    workbook.save(output_path)


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line_number, raw in enumerate(handle, 1):
            if not raw.strip():
                continue
            try:
                record = json.loads(raw)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{path}:{line_number}: invalid JSON: {exc}") from exc
            if not isinstance(record, dict):
                raise ValueError(f"{path}:{line_number}: each JSONL record must be an object")
            missing = {"case_id", "answer"} - record.keys()
            if missing:
                raise ValueError(f"{path}:{line_number}: missing fields {sorted(missing)}")
            if not isinstance(record["case_id"], str) or not record["case_id"]:
                raise ValueError(f"{path}:{line_number}: case_id must be a non-empty string")
            if not isinstance(record["answer"], str):
                raise ValueError(f"{path}:{line_number}: answer must be a string")
            records.append(record)

    if not records:
        raise ValueError(f"{path}: no records found")
    ids = [record["case_id"] for record in records]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{path}: duplicate case_id found; use one run per case per file")
    return records


def create_blind_sheet(
    cases_path: Path,
    run_paths: list[Path],
    output_path: Path,
    key_path: Path,
    seed: int = 20261013,
) -> int:
    """Create shuffled grading rows and a separate unblinding key."""
    if not run_paths:
        raise ValueError("At least one --runs file is required")
    if output_path.resolve() == key_path.resolve():
        raise ValueError("The grading sheet and unblinding key must be different files")

    cases = load_cases(cases_path)
    case_errors = validate_cases(cases)
    if case_errors:
        raise ValueError("Invalid benchmark cases:\n" + "\n".join(case_errors))
    cases_by_id = {case["id"]: case for case in cases}

    source_records: list[tuple[Path, dict[str, Any]]] = []
    expected_ids: set[str] | None = None
    for path in run_paths:
        records = _read_jsonl(path)
        ids = {record["case_id"] for record in records}
        unknown = ids - cases_by_id.keys()
        if unknown:
            raise ValueError(f"{path}: unknown case ids: {sorted(unknown)}")
        if expected_ids is None:
            expected_ids = ids
        elif ids != expected_ids:
            missing = sorted(expected_ids - ids)
            extra = sorted(ids - expected_ids)
            raise ValueError(
                f"Runs must contain the same cases for a fair comparison. "
                f"{path}: missing={missing}, extra={extra}"
            )
        source_records.extend((path, record) for record in records)

    if len(run_paths) < 2:
        raise ValueError("Blinded comparison requires at least two run files")

    rng = random.Random(seed)
    rng.shuffle(source_records)
    output_rows: list[dict[str, str]] = []
    key_rows: list[dict[str, str]] = []

    for index, (path, record) in enumerate(source_records, 1):
        grade_id = f"G-{index:04d}"
        case = cases_by_id[record["case_id"]]
        output_rows.append({
            "grade_id": grade_id,
            "case_id": case["id"],
            "category": str(case["category"]),
            "difficulty": str(case["difficulty"]),
            "prompt": str(case["prompt"]),
            "image_source": str(case.get("image_source", "")),
            "answer": _format_answer_for_review(record["answer"]),
            "reference_expectation": str(case["ground_truth"]),
            "trap_to_check": str(case["trap"]),
            "useful_next_evidence": str(case["next_evidence"]),
            **{dimension: "" for dimension in DIMENSIONS},
            "critical_failure": "",
            "critical_failure_reason": "",
            "grader_notes": "",
        })
        condition = str(record.get("condition", ""))
        prompt_id = str(record.get("prompt_id", ""))
        system_label = str(record.get("system", path.stem))
        # The runner's system field omits prompt version, so v0.1 and v0.2
        # would otherwise collapse into one aggregate. Keep prompt versions
        # distinct while preserving legacy labels where prompt_id is absent
        # or merely repeats the condition name.
        if condition == "investigator" and prompt_id and prompt_id != condition:
            if not system_label.endswith(f":{prompt_id}"):
                system_label = f"{system_label}:{prompt_id}"
        key_rows.append({
            "grade_id": grade_id,
            "case_id": case["id"],
            "source_file": str(path),
            "system": system_label,
            "provider": str(record.get("provider", "")),
            "model": str(record.get("model", "")),
            "condition": condition,
            "prompt_id": prompt_id,
            "prompt_sha256": str(record.get("prompt_sha256", "")),
            "latency_ms": str(record.get("latency_ms", "")),
        })

    for path in (output_path, key_path):
        path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.suffix.lower() == ".xlsx":
        _write_grading_workbook(output_path, output_rows)
    else:
        with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=GRADE_FIELDS)
            writer.writeheader()
            writer.writerows(output_rows)
    with key_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=KEY_FIELDS)
        writer.writeheader()
        writer.writerows(key_rows)

    return len(output_rows)


def _score_value(value: str, grade_id: str, dimension: str) -> int:
    try:
        score = int(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{grade_id}: {dimension} must be 0, 1, or 2") from exc
    if str(score) != value.strip() or score not in (0, 1, 2):
        raise ValueError(f"{grade_id}: {dimension} must be 0, 1, or 2")
    return score


def score_blind_sheet(
    grades_path: Path,
    key_path: Path,
    output_path: Path,
    summary_path: Path,
) -> dict[str, Any]:
    """Unblind a completed scoring sheet and write scored JSONL + metrics."""
    if grades_path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Reading .xlsx output requires openpyxl. Install with python -m pip install -e .") from exc
        workbook = load_workbook(grades_path, read_only=True, data_only=True)
        if "Review" not in workbook.sheetnames:
            raise ValueError("Workbook must contain a 'Review' sheet")
        sheet = workbook["Review"]
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            raise ValueError("Review sheet is empty")
        header_names = [str(value) if value is not None else "" for value in headers]
        grades = [
            {header: ("" if value is None else str(value)) for header, value in zip(header_names, values_row)}
            for values_row in values
        ]
        workbook.close()
    else:
        with grades_path.open(encoding="utf-8-sig", newline="") as handle:
            grades = list(csv.DictReader(handle))
    with key_path.open(encoding="utf-8-sig", newline="") as handle:
        key_rows = list(csv.DictReader(handle))

    key_by_id = {row["grade_id"]: row for row in key_rows}
    grade_ids = [row.get("grade_id", "") for row in grades]
    if not grades or len(grade_ids) != len(set(grade_ids)):
        raise ValueError("Scoring sheet is empty or contains duplicate grade_id values")
    if set(grade_ids) != set(key_by_id):
        raise ValueError("Scoring sheet and key must contain exactly the same grade_id values")

    scored_records: list[dict[str, Any]] = []
    by_system: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in grades:
        grade_id = row["grade_id"]
        key = key_by_id[grade_id]
        if row.get("case_id") != key.get("case_id"):
            raise ValueError(f"{grade_id}: case_id does not match the unblinding key")
        scores = {
            dimension: _score_value(row.get(dimension, ""), grade_id, dimension)
            for dimension in DIMENSIONS
        }
        critical_value = (row.get("critical_failure") or "").strip().lower()
        if critical_value not in {"yes", "no"}:
            raise ValueError(f"{grade_id}: critical_failure must be 'yes' or 'no'")
        failure_reason = (row.get("critical_failure_reason") or "").strip()
        if critical_value == "yes" and not failure_reason:
            raise ValueError(f"{grade_id}: explain the critical failure in critical_failure_reason")

        record = {
            "case_id": row["case_id"],
            "system": key["system"],
            "answer": row.get("answer", ""),
            "scores": scores,
            "critical_failure": critical_value == "yes",
            "critical_failure_reason": failure_reason,
            "grader_notes": (row.get("grader_notes") or "").strip(),
            "grade_id": grade_id,
            "provider": key.get("provider", ""),
            "model": key.get("model", ""),
            "condition": key.get("condition", ""),
            "prompt_id": key.get("prompt_id", ""),
            "source_file": key.get("source_file", ""),
            "latency_ms": key.get("latency_ms", ""),
        }
        record["score"] = score_record(record)
        scored_records.append(record)
        by_system[key["system"]].append(record)

    summary: dict[str, Any] = {
        "scored_answers": len(scored_records),
        "grading_note": "Weighted scores are comparison signals; critical failures are reported separately and must not be averaged away.",
        "systems": {},
    }
    for system, records in sorted(by_system.items()):
        dimension_means = {
            dim: round(sum(r["scores"][dim] for r in records) / len(records), 3)
            for dim in DIMENSIONS
        }
        failures = [
            {
                "case_id": r["case_id"],
                "grade_id": r["grade_id"],
                "reason": r["critical_failure_reason"],
            }
            for r in records if r["critical_failure"]
        ]
        latencies = sorted(
            float(r["latency_ms"]) for r in records if r.get("latency_ms", "") != ""
        )
        latency_metrics = {}
        if latencies:
            latency_metrics = {
                "mean_latency_ms": round(statistics.fmean(latencies), 1),
                "median_latency_ms": round(statistics.median(latencies), 1),
                "p95_latency_ms": round(latencies[min(len(latencies) - 1, math.ceil(0.95 * len(latencies)) - 1)], 1),
                "latency_samples": len(latencies),
            }
        summary["systems"][system] = {
            "cases": len(records),
            "mean_score": round(sum(r["score"] for r in records) / len(records), 2),
            "min_score": min(r["score"] for r in records),
            "max_score": max(r["score"] for r in records),
            "dimension_means_0_to_2": dimension_means,
            "critical_failures": len(failures),
            "critical_failure_rate": round(len(failures) / len(records), 4),
            "critical_failure_details": failures,
            **latency_metrics,
        }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        for record in scored_records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Create blinded Visual Investigator grading sheets.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    blind = subparsers.add_parser("blind", help="Create an anonymized CSV grading sheet.")
    blind.add_argument("--cases", type=Path, default=Path("benchmark/cases.json"))
    blind.add_argument("--runs", type=Path, nargs="+", required=True)
    blind.add_argument("--output", type=Path, required=True, help="Grading CSV output.")
    blind.add_argument("--key-output", type=Path, required=True, help="Keep this key private until grading is done.")
    blind.add_argument("--seed", type=int, default=20261013, help="Seed used to shuffle answer order.")

    score = subparsers.add_parser("score", help="Join completed ratings with the private key.")
    score.add_argument("--grades", type=Path, required=True, help="Completed CSV from the blind subcommand.")
    score.add_argument("--key", type=Path, required=True, help="Matching private key CSV.")
    score.add_argument("--output", type=Path, required=True, help="Scored JSONL output.")
    score.add_argument("--summary", type=Path, required=True, help="JSON summary output.")

    args = parser.parse_args()
    try:
        if args.command == "blind":
            count = create_blind_sheet(args.cases, args.runs, args.output, args.key_output, args.seed)
            print(f"Created blind grading sheet with {count} answers: {args.output}")
            print(f"Private unblinding key: {args.key_output}")
        else:
            summary = score_blind_sheet(args.grades, args.key, args.output, args.summary)
            print(f"Scored {summary['scored_answers']} answers.")
            print(f"Summary: {args.summary}")
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
